from __future__ import annotations

import argparse
import csv
import hashlib
import re
from pathlib import Path

from docx import Document


CARD_FIELDS = [
    "report_id",
    "work_id",
    "source_doc",
    "converted_docx",
    "sample_name",
    "manufacturer",
    "brand",
    "spec_model",
    "production_date",
    "sample_no",
    "test_type",
    "test_basis",
    "test_items",
    "test_date",
    "test_conclusion",
    "status",
    "error",
]

PARAMETER_FIELDS = [
    "report_id",
    "parameter_type",
    "numeric_value",
    "unit",
    "raw_value",
    "evidence",
    "source_kind",
    "source_index",
]

SUMMARY_FIELDS = ["section", "key", "count"]

META_ALIASES = {
    "sample_name": ("样品名称", "sample name", "产品名称"),
    "manufacturer": ("生产单位", "制造商", "manufacturer", "生产厂家", "厂家"),
    "brand": ("商标", "商 标", "brand"),
    "spec_model": ("规格型号", "型号规格", "规格", "型号", "type/model", "model"),
    "production_date": ("生产日期", "出厂日期", "manufacture date"),
    "sample_no": ("样品编号", "sample no", "样品号"),
    "test_type": ("检验类别", "检测类别", "test type"),
    "test_basis": ("检验依据", "检测依据", "test basis", "standard"),
    "test_items": ("检验项目", "检测项目", "test items"),
    "test_date": ("检验日期", "检测日期", "test date"),
    "test_conclusion": ("检验结论", "检测结论", "test conclusion"),
}

PARAMETER_PATTERNS = [
    ("thermal_transmittance_u", ("传热系数", "u值", "u 值", "u-value"), "W/(m2·K)"),
    ("thermal_resistance_r", ("热阻", "传热阻"), "m2·K/W"),
    ("thermal_conductivity_lambda", ("导热系数",), "W/(m·K)"),
]


def clean(value: object) -> str:
    return " ".join(str(value or "").replace("\u3000", " ").split()).strip()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def safe_stem(value: str) -> str:
    text = re.sub(r'[<>:"/\\|*?]+', "_", clean(value))
    return text[:120] or "report"


def resolve_doc_path(row: dict[str, str], data_root: Path) -> Path:
    for key in ("destination", "source_path"):
        value = clean(row.get(key))
        if not value:
            continue
        path = Path(value)
        if path.is_absolute() and path.exists():
            return path
        candidate = data_root / value
        if candidate.exists():
            return candidate
    return data_root / clean(row.get("destination") or row.get("source_path"))


def docx_cache_path(source: Path, report_id: str, cache_dir: Path) -> Path:
    digest = hashlib.sha1(str(source).encode("utf-8")).hexdigest()[:10]
    return cache_dir / f"{safe_stem(report_id or source.stem)}__{digest}.docx"


def convert_to_docx(source: Path, target: Path, word_app=None) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        return target
    if source.suffix.lower() == ".docx":
        import shutil

        shutil.copy2(source, target)
        return target
    if source.suffix.lower() != ".doc":
        raise ValueError(f"Unsupported Word file type: {source.suffix}")
    if word_app is None:
        raise ValueError("word_app is required for .doc conversion")
    doc = None
    try:
        doc = word_app.Documents.Open(str(source.resolve()), ReadOnly=True, AddToRecentFiles=False)
        doc.SaveAs(str(target.resolve()), FileFormat=16)
        doc.Close(False)
        return target
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass


def unique_row_cells(row) -> list[str]:
    cells: list[str] = []
    seen = set()
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id in seen:
            continue
        seen.add(tc_id)
        text = clean(cell.text)
        if text:
            cells.append(text)
    return cells


def field_key(field: str) -> str | None:
    text = clean(field).lower().replace(" ", "")
    for key, aliases in META_ALIASES.items():
        for alias in aliases:
            if alias.lower().replace(" ", "") in text:
                return key
    return None


def append_unique(existing: str, value: str) -> str:
    existing = clean(existing)
    value = clean(value)
    if not value:
        return existing
    if not existing:
        return value
    parts = [part.strip() for part in existing.split(";") if part.strip()]
    if value not in parts:
        parts.append(value)
    return "; ".join(parts)


def fill_card_from_pair(card: dict[str, str], field: str, value: str) -> None:
    key = field_key(field)
    if not key:
        return
    if key in {"test_items", "test_basis", "test_conclusion"}:
        card[key] = append_unique(card[key], value)
    elif not card.get(key):
        card[key] = clean(value)


def extract_card(doc: Document, base: dict[str, str]) -> dict[str, str]:
    card = {field: "" for field in CARD_FIELDS}
    card.update(base)
    for table in doc.tables[:5]:
        for row in table.rows[:25]:
            cells = unique_row_cells(row)
            if len(cells) >= 4:
                for idx in range(0, len(cells) - 1, 2):
                    fill_card_from_pair(card, cells[idx], cells[idx + 1])
            elif len(cells) == 2:
                fill_card_from_pair(card, cells[0], cells[1])
    for paragraph in doc.paragraphs[:80]:
        text = clean(paragraph.text)
        if "检验结论" in text or "检测结论" in text:
            card["test_conclusion"] = append_unique(card["test_conclusion"], text)
    return card


THERMAL_PATTERNS = [
    (
        "thermal_transmittance_u",
        re.compile(
            r"(?:传热系数|K\s*[=：:])[^。；;\n]{0,40}?"
            r"([-+]?\d+(?:\.\d+)?)\s*"
            r"(W\s*/\s*[（(]?\s*m[²2]?\s*[·.]?\s*K\s*[）)]?)",
            re.I,
        ),
    ),
    (
        "thermal_resistance_r",
        re.compile(
            r"(?:热阻|R\s*[=：:])[^。；;\n]{0,40}?"
            r"([-+]?\d+(?:\.\d+)?)\s*"
            r"([（(]?\s*m[²2]\s*[·.]?\s*K\s*[）)]?\s*/\s*W)",
            re.I,
        ),
    ),
    (
        "thermal_conductivity_lambda",
        re.compile(
            r"(?:导热系数)[^。；;\n]{0,60}?"
            r"([-+]?\d+(?:\.\d+)?)\s*"
            r"(W\s*/\s*[（(]?\s*m\s*[·.]?\s*K\s*[）)]?)",
            re.I,
        ),
    ),
]


def parameter_type_for_row(text: str) -> tuple[str, str]:
    lowered = clean(text).lower()
    for parameter_type, keywords, default_unit in PARAMETER_PATTERNS:
        if any(keyword.lower() in lowered for keyword in keywords):
            return parameter_type, default_unit
    return "", ""


def parse_value_cell(text: str, default_unit: str) -> tuple[str, str]:
    value = clean(text)
    number_match = re.search(r"[-+]?\d+(?:\.\d+)?", value)
    number = number_match.group(0) if number_match else ""
    unit = ""
    unit_patterns = [
        r"W/[（(]?m[²2]?[·.]?K[）)]?",
        r"W/\(?m[²2]?[·.]?K\)?",
        r"[（(]?m[²2][·.]?K[）)]?/W",
        r"\(?m[²2][·.]?K\)?/W",
    ]
    for pattern in unit_patterns:
        match = re.search(pattern, value, re.I)
        if match:
            unit = clean(match.group(0))
            break
    return number, unit or default_unit


def measured_value_cell(cells: list[str]) -> str:
    for cell in reversed(cells):
        text = clean(cell)
        if text in {"符合", "不符合", "合格", "不合格", "——", "-", "/"}:
            continue
        if re.search(r"\d", text) and re.search(r"(W|m[²2]|K)", text, re.I):
            return text
    return ""


def extract_thermal_mentions(text: str) -> list[dict[str, str]]:
    value = clean(text)
    mentions = []
    for parameter_type, pattern in THERMAL_PATTERNS:
        for match in pattern.finditer(value):
            start = max(match.start() - 24, 0)
            end = min(match.end() + 24, len(value))
            mentions.append(
                {
                    "parameter_type": parameter_type,
                    "numeric_value": match.group(1),
                    "unit": clean(match.group(2)),
                    "raw_value": clean(match.group(0)),
                    "evidence": clean(value[start:end]),
                }
            )
    return mentions


def dedupe_parameters(rows: list[dict[str, str]]) -> list[dict[str, str]]:
    seen = set()
    out = []
    for row in rows:
        key = (
            row["report_id"],
            row["parameter_type"],
            row["numeric_value"],
            row["unit"],
            row["evidence"],
        )
        if key in seen:
            continue
        seen.add(key)
        out.append(row)
    return out


def extract_parameters(doc: Document, report_id: str) -> list[dict[str, str]]:
    params = []
    for table_index, table in enumerate(doc.tables):
        for row_index, row in enumerate(table.rows):
            cells = unique_row_cells(row)
            if not cells:
                continue
            row_text = " / ".join(cells)
            row_parameter_type, default_unit = parameter_type_for_row(row_text)
            value_cell = measured_value_cell(cells) if len(cells) >= 4 else ""
            if row_parameter_type and value_cell:
                number, unit = parse_value_cell(value_cell, default_unit)
                if number:
                    params.append(
                        {
                            "report_id": report_id,
                            "parameter_type": row_parameter_type,
                            "numeric_value": number,
                            "unit": unit,
                            "raw_value": value_cell,
                            "evidence": row_text,
                            "source_kind": "table_row",
                            "source_index": f"{table_index}:{row_index}",
                        }
                    )
                continue
            for mention in extract_thermal_mentions(row_text):
                params.append(
                    {
                        "report_id": report_id,
                        **mention,
                        "source_kind": "table_row",
                        "source_index": f"{table_index}:{row_index}",
                    }
                )

    for index, paragraph in enumerate(doc.paragraphs):
        text = clean(paragraph.text)
        if not text:
            continue
        for mention in extract_thermal_mentions(text):
            params.append(
                {
                    "report_id": report_id,
                    **mention,
                    "source_kind": "paragraph",
                    "source_index": str(index),
                }
            )
    return dedupe_parameters(params)


def start_word():
    import win32com.client as win32

    word = win32.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    return word


def stop_word(word) -> None:
    try:
        word.Quit()
    except Exception:
        pass


def write_summary(path: Path, cards: list[dict[str, str]], params: list[dict[str, str]]) -> None:
    rows = [
        {"section": "total", "key": "report_cards", "count": str(len(cards))},
        {"section": "total", "key": "thermal_parameters", "count": str(len(params))},
    ]
    for status in sorted({row["status"] for row in cards}):
        rows.append(
            {
                "section": "cards_by_status",
                "key": status,
                "count": str(sum(1 for row in cards if row["status"] == status)),
            }
        )
    for parameter_type in sorted({row["parameter_type"] for row in params}):
        rows.append(
            {
                "section": "parameter_type",
                "key": parameter_type,
                "count": str(sum(1 for row in params if row["parameter_type"] == parameter_type)),
            }
        )
    write_csv(path, rows, SUMMARY_FIELDS)


def run(args: argparse.Namespace) -> tuple[list[dict[str, str]], list[dict[str, str]]]:
    manifest_rows = read_csv(Path(args.manifest))
    selected = manifest_rows[args.offset : args.offset + args.limit if args.limit else None]
    data_root = Path(args.data_root)
    output_dir = Path(args.output_dir)
    cache_dir = output_dir / "docx_cache"
    cards = []
    params = []

    word = start_word()
    try:
        for row in selected:
            report_id = clean(row.get("report_id"))
            source = resolve_doc_path(row, data_root)
            target = docx_cache_path(source, report_id, cache_dir)
            base = {
                "report_id": report_id,
                "work_id": clean(row.get("work_id")),
                "source_doc": str(source),
                "converted_docx": str(target),
                "status": "OK",
                "error": "",
            }
            try:
                convert_to_docx(source, target, word)
                doc = Document(str(target))
                card = extract_card(doc, base)
                cards.append(card)
                params.extend(extract_parameters(doc, report_id))
            except Exception as exc:
                base["status"] = "FAIL"
                base["error"] = repr(exc)
                cards.append({field: base.get(field, "") for field in CARD_FIELDS})
    finally:
        stop_word(word)

    return cards, params


def main() -> None:
    parser = argparse.ArgumentParser(description="Extract wall-insulation report cards and thermal parameters directly from Word reports.")
    parser.add_argument(
        "--manifest",
        default=str(Path(__file__).resolve().parents[1] / "data" / "fine_classification" / "wall_insulation.csv"),
    )
    parser.add_argument("--data-root", default=r"G:\shujuku")
    parser.add_argument(
        "--output-dir",
        default=str(Path(__file__).resolve().parents[1] / "data" / "wall_insulation_word_extraction"),
    )
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=50)
    args = parser.parse_args()

    cards, params = run(args)
    output_dir = Path(args.output_dir)
    write_csv(output_dir / "report_cards.csv", cards, CARD_FIELDS)
    write_csv(output_dir / "thermal_parameters.csv", params, PARAMETER_FIELDS)
    write_summary(output_dir / "summary.csv", cards, params)
    print(f"Report cards: {len(cards)}")
    print(f"Thermal parameters: {len(params)}")
    print(f"Output: {output_dir}")


if __name__ == "__main__":
    main()
