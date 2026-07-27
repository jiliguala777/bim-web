from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document


TRACE_FIELDS = [
    "report_id",
    "sequence",
    "block_type",
    "table_index",
    "row_index",
    "col_index",
    "text",
    "row_text",
]
META_FIELDS = ["report_id", "table_index", "row_index", "field", "value"]
ITEM_FIELDS = [
    "report_id",
    "table_index",
    "row_index",
    "section",
    "seq",
    "item_name",
    "requirement",
    "result",
    "conclusion",
    "raw_row",
]
CATEGORY_FIELDS = ["report_id", "category", "load_relevant", "reason"]


@dataclass
class ExtractResult:
    report_id: str
    raw_blocks: list[dict[str, str]]
    meta_rows: list[dict[str, str]]
    item_rows: list[dict[str, str]]
    category_rows: list[dict[str, str]]


HEADER_ALIASES = {
    "seq": ("序号", "no.", "no"),
    "item_name": ("检验项目", "检测项目", "test items", "项目名称"),
    "requirement": ("技术要求", "requirement of standard", "标准要求"),
    "result": ("检测结果", "检验结果", "test results", "结果"),
    "conclusion": ("单项判定", "item conclusion", "判定", "结论"),
}

FIELD_HINTS = (
    "编号",
    "日期",
    "名称",
    "型号",
    "规格",
    "单位",
    "依据",
    "地点",
    "类别",
    "等级",
    "数量",
    "状态",
    "样品",
    "委托",
    "生产",
    "检测",
    "检验",
    "报告",
    "sample",
    "client",
    "test type",
    "address",
    "brand",
    "manufacturer",
    "type/model",
    "model",
    "quantity",
    "date",
    "commission",
    "test place",
    "standard",
    "criteria",
    "reference",
)


def clean_text(value: object) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("\u3000", " ").split()).strip()


def _contains_any(text: str, needles: Iterable[str]) -> bool:
    lowered = text.lower()
    return any(needle.lower() in lowered for needle in needles)


def _is_blank(value: str) -> bool:
    return clean_text(value) in {"", "-", "--", "------", "—", "/"}


def _looks_like_field(value: str) -> bool:
    text = clean_text(value).strip(":：")
    if not text or len(text) > 50:
        return False
    if re.fullmatch(r"[\d.]+", text):
        return False
    return _contains_any(text, FIELD_HINTS) or text.endswith((":", "："))


def _looks_like_item_header(cells: list[str]) -> bool:
    joined = " ".join(cells).lower()
    return (
        _contains_any(joined, HEADER_ALIASES["seq"])
        and _contains_any(joined, HEADER_ALIASES["item_name"])
        and (
            _contains_any(joined, HEADER_ALIASES["result"])
            or _contains_any(joined, HEADER_ALIASES["requirement"])
        )
    )


def _header_role(cell: str) -> str | None:
    text = clean_text(cell).lower()
    for role, aliases in HEADER_ALIASES.items():
        if any(alias.lower() in text for alias in aliases):
            return role
    return None


def _dedupe_cells(cells: list[str]) -> list[str]:
    return [value for cell in cells if (value := clean_text(cell))]


def _unique_row_cells(row) -> list[tuple[int, str]]:
    cells: list[tuple[int, str]] = []
    seen_tc_ids: set[int] = set()
    for col_index, cell in enumerate(row.cells):
        tc_id = id(cell._tc)
        if tc_id in seen_tc_ids:
            continue
        seen_tc_ids.add(tc_id)
        cells.append((col_index, clean_text(cell.text)))
    return cells


def _table_rows(table) -> list[list[str]]:
    return [[text for _, text in _unique_row_cells(row)] for row in table.rows]


def extract_meta_from_rows(report_id: str, table_index: int, rows: list[list[str]]) -> list[dict[str, str]]:
    meta_rows: list[dict[str, str]] = []
    seen: set[tuple[int, int, str, str]] = set()

    for row_index, row in enumerate(rows):
        cells = _dedupe_cells(row)
        if not cells or _looks_like_item_header(cells):
            continue

        pairs: list[tuple[str, str]] = []
        if len(cells) >= 4:
            for i in range(0, len(cells) - 1, 2):
                field, value = cells[i], cells[i + 1]
                if _looks_like_field(field) and not _is_blank(value):
                    pairs.append((field.strip(":："), value))

        if not pairs and len(cells) == 2:
            field, value = cells
            if _looks_like_field(field) and not _is_blank(value):
                pairs.append((field.strip(":："), value))

        for field, value in pairs:
            key = (table_index, row_index, field, value)
            if key in seen:
                continue
            seen.add(key)
            meta_rows.append(
                {
                    "report_id": report_id,
                    "table_index": str(table_index),
                    "row_index": str(row_index),
                    "field": field,
                    "value": value,
                }
            )

    return meta_rows


def extract_items_from_rows(report_id: str, table_index: int, rows: list[list[str]]) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    header_map: dict[int, str] | None = None
    section = ""

    for row_index, row in enumerate(rows):
        cells = _dedupe_cells(row)
        if not cells:
            continue

        if len(cells) == 1:
            section = cells[0]
            continue

        if _looks_like_item_header(cells):
            header_map = {}
            for col_index, cell in enumerate(cells):
                role = _header_role(cell)
                if role:
                    header_map[col_index] = role
            continue

        if not header_map:
            continue

        first = cells[0] if cells else ""
        if not re.fullmatch(r"\d{1,4}([.)、])?", first):
            continue

        record = {
            "report_id": report_id,
            "table_index": str(table_index),
            "row_index": str(row_index),
            "section": section,
            "seq": "",
            "item_name": "",
            "requirement": "",
            "result": "",
            "conclusion": "",
            "raw_row": " | ".join(cells),
        }

        if len(cells) > len(header_map) and len(cells) >= 5:
            record["seq"] = cells[0]
            record["conclusion"] = cells[-1]
            record["result"] = cells[-2]
            record["requirement"] = cells[-3]
            record["item_name"] = " / ".join(dict.fromkeys(cells[1:-3]))
        else:
            for col_index, role in header_map.items():
                if col_index < len(cells):
                    record[role] = cells[col_index]
        items.append(record)

    return items


def extract_docx_content(docx_path: str | Path, category: str = "wall_insulation") -> ExtractResult:
    path = Path(docx_path)
    report_id = path.stem
    doc = Document(str(path))

    raw_blocks: list[dict[str, str]] = []
    sequence = 0

    for paragraph in doc.paragraphs:
        text = clean_text(paragraph.text)
        if not text:
            continue
        raw_blocks.append(
            {
                "report_id": report_id,
                "sequence": str(sequence),
                "block_type": "paragraph",
                "table_index": "",
                "row_index": "",
                "col_index": "",
                "text": text,
                "row_text": "",
            }
        )
        sequence += 1

    meta_rows: list[dict[str, str]] = []
    item_rows: list[dict[str, str]] = []

    for table_index, table in enumerate(doc.tables):
        rows = _table_rows(table)
        for row_index, row in enumerate(rows):
            row_text = " | ".join([cell for cell in row if cell])
            if row_text:
                raw_blocks.append(
                    {
                        "report_id": report_id,
                        "sequence": str(sequence),
                        "block_type": "table_row",
                        "table_index": str(table_index),
                        "row_index": str(row_index),
                        "col_index": "",
                        "text": row_text,
                        "row_text": row_text,
                    }
                )
                sequence += 1

            for col_index, cell_text in _unique_row_cells(table.rows[row_index]):
                if not cell_text:
                    continue
                raw_blocks.append(
                    {
                        "report_id": report_id,
                        "sequence": str(sequence),
                        "block_type": "table_cell",
                        "table_index": str(table_index),
                        "row_index": str(row_index),
                        "col_index": str(col_index),
                        "text": cell_text,
                        "row_text": row_text,
                    }
                )
                sequence += 1

        meta_rows.extend(extract_meta_from_rows(report_id, table_index, rows))
        item_rows.extend(extract_items_from_rows(report_id, table_index, rows))

    return ExtractResult(
        report_id=report_id,
        raw_blocks=raw_blocks,
        meta_rows=meta_rows,
        item_rows=item_rows,
        category_rows=[
            {
                "report_id": report_id,
                "category": category,
                "load_relevant": "yes",
                "reason": f"fine_classified_{category}",
            }
        ],
    )


def write_csv(path: Path, fieldnames: list[str], rows: list[dict[str, str]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def write_report_output(result: ExtractResult, output_root: str | Path) -> Path:
    report_dir = Path(output_root) / result.report_id
    report_dir.mkdir(parents=True, exist_ok=True)
    write_csv(report_dir / "trace.csv", TRACE_FIELDS, result.raw_blocks)
    write_csv(report_dir / "meta.csv", META_FIELDS, result.meta_rows)
    write_csv(report_dir / "items.csv", ITEM_FIELDS, result.item_rows)
    write_csv(report_dir / "category.csv", CATEGORY_FIELDS, result.category_rows)
    return report_dir
