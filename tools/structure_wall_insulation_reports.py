from __future__ import annotations

import argparse
import csv
import hashlib
import shutil
from pathlib import Path

from docx import Document

from report_extractor import extract_docx_content, write_report_output


BASE_DIR = Path(__file__).resolve().parents[1]
DEFAULT_MANIFEST = BASE_DIR / "data" / "fine_classification" / "wall_insulation.csv"
DEFAULT_OUTPUT_ROOT = BASE_DIR / "data" / "wall_insulation_structured_full"
DEFAULT_DOCX_ROOT = BASE_DIR / "data" / "wall_insulation_structured_docx"
DEFAULT_CATEGORY = "wall_insulation"
REPORT_OUTPUT_FILES = ("trace.csv", "meta.csv", "items.csv", "category.csv")

SUMMARY_FIELDS = [
    "report_id",
    "status",
    "source_doc",
    "converted_docx",
    "raw_blocks",
    "meta_rows",
    "item_rows",
    "category",
    "load_relevant",
    "error",
    "source_path",
    "destination",
    "work_id",
    "structure_type",
    "table_count",
    "paragraph_count",
    "char_count",
    "item_header_rows",
    "item_body_rows",
]


def clean(value: object) -> str:
    return " ".join(str(value or "").replace("\u3000", " ").split()).strip()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def write_csv(path: Path, rows: list[dict[str, str]], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def append_csv(path: Path, row: dict[str, str], fields: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    write_header = not path.exists() or path.stat().st_size == 0
    with path.open("a", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        if write_header:
            writer.writeheader()
        writer.writerow(row)


def resolve_doc_path(row: dict[str, str], data_root: Path) -> Path:
    for key in ("destination", "source_path"):
        value = clean(row.get(key))
        if not value:
            continue
        path = Path(value)
        if path.is_absolute() and path.exists():
            return path
        candidate = data_root / value
        if candidate.exists():
            return candidate
    return data_root / clean(row.get("destination") or row.get("source_path"))


def safe_stem(value: str) -> str:
    cleaned = "".join("_" if ch in '<>:"/\\|*?' else ch for ch in clean(value))
    return cleaned[:120] or "report"


def docx_output_path(source: Path, report_id: str, docx_root: Path) -> Path:
    digest = hashlib.sha1(str(source).encode("utf-8")).hexdigest()[:12]
    return docx_root / f"{safe_stem(report_id or source.stem)}__{digest}.docx"


def convert_doc_to_docx(source: Path, target: Path, word_app=None) -> Path:
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists():
        return target
    if source.suffix.lower() == ".docx":
        shutil.copy2(source, target)
        return target
    if source.suffix.lower() != ".doc":
        raise ValueError(f"Unsupported Word file type: {source.suffix}")
    if word_app is None:
        raise ValueError("word_app is required for .doc conversion")
    doc = None
    try:
        doc = word_app.Documents.Open(str(source.resolve()), ReadOnly=True, AddToRecentFiles=False)
        doc.SaveAs(str(target.resolve()), FileFormat=16)
        doc.Close(False)
        return target
    finally:
        if doc is not None:
            try:
                doc.Close(False)
            except Exception:
                pass


def _dedupe_cells(row) -> list[str]:
    cells: list[str] = []
    seen_tc_ids: set[int] = set()
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id in seen_tc_ids:
            continue
        seen_tc_ids.add(tc_id)
        text = clean(cell.text)
        if text:
            cells.append(text)
    return cells


def _is_seq(value: str) -> bool:
    text = clean(value).rstrip(".、)")
    return text.isdigit() and 1 <= len(text) <= 4


def _looks_like_item_header(cells: list[str]) -> bool:
    text = " ".join(cells).lower()
    return (
        "序号" in text
        and ("检验项目" in text or "检测项目" in text or "test items" in text)
        and ("检测结果" in text or "检验结果" in text or "test results" in text or "结果" in text)
        and ("技术要求" in text or "requirement" in text or "标准要求" in text)
    )


def detect_table_profile(docx_path: Path) -> dict[str, str]:
    doc = Document(str(docx_path))
    paragraph_count = sum(1 for paragraph in doc.paragraphs if clean(paragraph.text))
    char_count = sum(len(clean(paragraph.text)) for paragraph in doc.paragraphs if clean(paragraph.text))
    item_header_rows = 0
    item_body_rows = 0
    for table in doc.tables:
        in_items = False
        for row in table.rows:
            cells = _dedupe_cells(row)
            if not cells:
                continue
            if _looks_like_item_header(cells):
                item_header_rows += 1
                in_items = True
                continue
            if in_items and _is_seq(cells[0]):
                item_body_rows += 1
                continue
            if in_items and not _is_seq(cells[0]):
                in_items = False
    if item_header_rows and item_body_rows:
        structure_type = "TABLE_OK"
    elif doc.tables and paragraph_count > 2000:
        structure_type = "MIXED"
    elif doc.tables:
        structure_type = "TABLE_WEAK"
    elif paragraph_count or char_count:
        structure_type = "TEXT"
    else:
        structure_type = "UNKNOWN"
    return {
        "structure_type": structure_type,
        "table_count": str(len(doc.tables)),
        "paragraph_count": str(paragraph_count),
        "char_count": str(char_count),
        "item_header_rows": str(item_header_rows),
        "item_body_rows": str(item_body_rows),
    }


def is_report_complete(output_root: Path, report_id: str) -> bool:
    report_dir = output_root / report_id
    return all((report_dir / name).exists() for name in REPORT_OUTPUT_FILES)


def start_word():
    import win32com.client as win32

    word = win32.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    return word


def stop_word(word) -> None:
    try:
        word.Quit()
    except Exception:
        pass


def summary_row(
    row: dict[str, str],
    source_doc: Path,
    converted_docx: Path,
    status: str,
    profile: dict[str, str],
    result=None,
    error: str = "",
) -> dict[str, str]:
    category_row = result.category_rows[0] if result and result.category_rows else {}
    return {
        "report_id": result.report_id if result else clean(row.get("report_id")) or converted_docx.stem,
        "status": status,
        "source_doc": str(source_doc),
        "converted_docx": str(converted_docx),
        "raw_blocks": str(len(result.raw_blocks)) if result else "0",
        "meta_rows": str(len(result.meta_rows)) if result else "0",
        "item_rows": str(len(result.item_rows)) if result else "0",
        "category": category_row.get("category", "wall_insulation"),
        "load_relevant": category_row.get("load_relevant", "yes"),
        "error": error,
        "source_path": clean(row.get("source_path")),
        "destination": clean(row.get("destination")),
        "work_id": clean(row.get("work_id")),
        **profile,
    }


def run(args: argparse.Namespace) -> list[dict[str, str]]:
    manifest_rows = read_csv(Path(args.manifest))
    selected = manifest_rows[args.offset : args.offset + args.limit if args.limit else None]
    data_root = Path(args.data_root)
    output_root = Path(args.output_root)
    docx_root = Path(args.docx_root)
    summary_path = output_root / "_summary.csv"
    output_root.mkdir(parents=True, exist_ok=True)

    if summary_path.exists() and not args.resume:
        summary_path.unlink()

    rows: list[dict[str, str]] = []
    word = start_word()
    try:
        for manifest_row in selected:
            source_doc = resolve_doc_path(manifest_row, data_root)
            converted_docx = docx_output_path(source_doc, clean(manifest_row.get("report_id")), docx_root)
            try:
                if args.resume and converted_docx.exists() and is_report_complete(output_root, converted_docx.stem):
                    continue
                convert_doc_to_docx(source_doc, converted_docx, word)
                profile = detect_table_profile(converted_docx)
                result = extract_docx_content(converted_docx, category=args.category)

                # Keep the useful trace/meta for every report. For weak tables, preserve items
                # when the generic parser finds them instead of intentionally blanking them.
                write_report_output(result, output_root)
                row = summary_row(manifest_row, source_doc, converted_docx, profile["structure_type"], profile, result)
            except Exception as exc:
                profile = {
                    "structure_type": "FAIL",
                    "table_count": "0",
                    "paragraph_count": "0",
                    "char_count": "0",
                    "item_header_rows": "0",
                    "item_body_rows": "0",
                }
                row = summary_row(manifest_row, source_doc, converted_docx, "FAIL", profile, error=repr(exc))
            rows.append(row)
            append_csv(summary_path, row, SUMMARY_FIELDS)
    finally:
        stop_word(word)
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description="Structure fine-classified Word reports into CSV report folders.")
    parser.add_argument("--manifest", default=str(DEFAULT_MANIFEST))
    parser.add_argument("--data-root", default=r"G:\shujuku")
    parser.add_argument("--output-root", default=str(DEFAULT_OUTPUT_ROOT))
    parser.add_argument("--docx-root", default=str(DEFAULT_DOCX_ROOT))
    parser.add_argument("--category", default=DEFAULT_CATEGORY)
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=50)
    parser.add_argument("--resume", action="store_true")
    args = parser.parse_args()

    rows = run(args)
    counts: dict[str, int] = {}
    for row in rows:
        counts[row["status"]] = counts.get(row["status"], 0) + 1
    print(f"Structured {args.category} reports: total={len(rows)} counts={counts}")
    print(f"Output: {Path(args.output_root)}")


if __name__ == "__main__":
    main()
